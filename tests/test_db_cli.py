import io
import json
import sqlite3
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "vendor"))
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from playauto.db_cli import main as db_main
from tests.test_db import make_crawl
from tests.test_db_export import make_reference


class DbCliTests(unittest.TestCase):
    def test_import_then_edit_then_export(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            db = tmp / "playauto.db"
            crawl = tmp / "crawl.xlsx"
            reference = tmp / "reference.xlsx"
            out = tmp / "out"
            make_crawl(
                crawl,
                rows=[[
                    1, "https://img/1.jpg", "상품1", "7,080원", "https://order/1",
                    None, None, None, None, None, None, None, None, None, None, None, None,
                ]],
            )
            make_reference(reference)

            stdout = io.StringIO()
            with redirect_stdout(stdout):
                code = db_main([
                    "db-import", "--db", str(db), "--crawl", str(crawl),
                    "--category-code", "기타재화[35]",
                ])
            self.assertEqual(code, 0)
            self.assertEqual(json.loads(stdout.getvalue())["inserted"], 1)

            conn = sqlite3.connect(db)
            conn.execute("UPDATE products SET category_code='가공식품[21]', detail='수정상세'")
            conn.execute("UPDATE listings SET margin_rate=0.2, sale_name='수정상품'")
            conn.commit()
            conn.close()

            stdout = io.StringIO()
            with redirect_stdout(stdout):
                code = db_main([
                    "db-export", "--db", str(db), "--reference", str(reference),
                    "--output-dir", str(out), "--date", "2026-10-27",
                ])
            result = json.loads(stdout.getvalue())
            setting = load_workbook(result["setting"])["setting"]
            single = load_workbook(result["single"]).active

            out_coupang = tmp / "out_coupang"
            stdout = io.StringIO()
            with redirect_stdout(stdout):
                code_coupang = db_main([
                    "db-export", "--db", str(db), "--reference", str(reference),
                    "--output-dir", str(out_coupang), "--date", "2026-10-27",
                    "--coupang-single",
                ])
            result_coupang = json.loads(stdout.getvalue())
            coupang_ws = load_workbook(result_coupang["single"]).active
            coupang_single_exists = Path(result_coupang["single"]).exists()

        self.assertEqual(code, 0)
        self.assertEqual(setting["A2"].value, "가공식품[21]")
        self.assertEqual(setting["D2"].value, "수정상품")
        self.assertEqual(setting["F2"].value, 0.2)
        self.assertEqual(setting["H2"].value, "수정상세")
        self.assertEqual(setting["R2"].value, "https://order/1")
        self.assertNotIn("주문링크", [cell.value for cell in single[1]])
        self.assertEqual(code_coupang, 0)
        self.assertTrue(coupang_single_exists)
        self.assertIsNone(result_coupang["general"])
        self.assertNotIn("주문링크", [cell.value for cell in coupang_ws[1]])
