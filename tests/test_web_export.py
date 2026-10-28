import hashlib
import io
import sqlite3
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from playauto.db import connect, ensure_schema
from playauto.web.export_service import (
    delete_export,
    export_to_dir,
    list_exports,
    save_reference,
    validate_reference,
)
from playauto.web.queries import get_setting, set_setting
from tests.web_helpers import (
    login,
    make_web_app,
    post_with_csrf,
    seed_product,
)


def make_reference(path):
    wb = Workbook()
    setting = wb.active
    setting.title = "setting"
    setting.append(["old"])
    template = wb.create_sheet("template")
    template.append(
        [
            "쇼핑몰",
            "계정",
            "템플릿코드정보",
            "템플릿코드",
            "머리말/꼬리말 템플릿코드",
            "단일/일반",
        ]
    )
    template.append(["SmartStore", "A", "1박스", "111", "900", "일반"])
    wb.save(path)
    wb.close()


class ExportServiceTests(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.reference = self.root / "current.xlsx"
        make_reference(self.reference)
        self.app, self.db_path = make_web_app(self.tmp.name)
        conn = connect(self.db_path)
        ensure_schema(conn)
        seed_product(conn)
        conn.close()

    def test_validate_reference(self):
        validate_reference(self.reference)
        bad = self.root / "bad.xlsx"
        wb = Workbook()
        wb.active.title = "setting"
        wb.save(bad)
        wb.close()
        with self.assertRaises(ValueError):
            validate_reference(bad)

    def test_save_reference_replaces_and_hashes(self):
        src = self.root / "upload.xlsx"
        make_reference(src)
        dest = self.root / "stored" / "current.xlsx"
        info = save_reference(src, dest)
        self.assertTrue(dest.exists())
        self.assertEqual(
            info["sha256"],
            hashlib.sha256(dest.read_bytes()).hexdigest(),
        )

    def test_export_creates_files_and_list_delete(self):
        result = export_to_dir(
            self.db_path,
            self.reference,
            self.root / "exports",
            option="기타재화[35]",
            run_date=None,
            coupang_single=False,
        )
        self.assertTrue(result.setting.exists())
        self.assertTrue(result.single.exists())
        self.assertTrue(result.general.exists())
        wb = load_workbook(result.setting)
        try:
            self.assertIn("주문링크", [c.value for c in wb["setting"][1]])
        finally:
            wb.close()
        exports = list_exports(self.root / "exports")
        self.assertEqual(len(exports), 1)
        self.assertTrue(delete_export(self.root / "exports", result.export_id))
        self.assertFalse(result.directory.exists())

    def test_export_missing_reference_raises(self):
        with self.assertRaises(ValueError):
            export_to_dir(
                self.db_path,
                self.root / "missing.xlsx",
                self.root / "exports",
                option="기타재화[35]",
            )


class ExportRouteTests(unittest.TestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.reference = self.root / "current.xlsx"
        make_reference(self.reference)
        self.app, self.db_path = make_web_app(self.tmp.name)
        self.client = self.app.test_client()
        login(self.client)
        self.conn = connect(self.db_path)
        ensure_schema(self.conn)
        self.addCleanup(self.conn.close)
        seed_product(self.conn)

    def upload(self, path, name="새 기준.xlsx"):
        payload = Path(path).read_bytes()
        return post_with_csrf(
            self.client,
            "/exports/reference",
            {"file": (io.BytesIO(payload), name)},
        )

    def make_changed_reference(self):
        source = self.root / "upload.xlsx"
        make_reference(source)
        wb = load_workbook(source)
        try:
            wb["setting"].append(["changed"])
            wb.save(source)
        finally:
            wb.close()
        return source

    def test_reference_upload_records_metadata(self):
        source = self.make_changed_reference()
        payload = source.read_bytes()

        response = self.upload(source)

        self.assertEqual(response.status_code, 302)
        self.assertEqual(self.reference.read_bytes(), payload)
        self.assertEqual(
            get_setting(self.conn, "reference_original_name"),
            "새 기준.xlsx",
        )
        self.assertEqual(
            get_setting(self.conn, "reference_sha256"),
            hashlib.sha256(payload).hexdigest(),
        )

    def test_reference_upload_requires_csrf(self):
        source = self.make_changed_reference()

        response = self.client.post(
            "/exports/reference",
            data={"file": (io.BytesIO(source.read_bytes()), "x.xlsx")},
        )

        self.assertEqual(response.status_code, 400)

    def test_reference_upload_restores_previous_on_settings_failure(self):
        previous = self.reference.read_bytes()
        source = self.make_changed_reference()

        with patch(
            "playauto.web.routes.set_settings",
            side_effect=sqlite3.OperationalError("boom"),
        ), patch.object(self.app.logger, "exception") as logged:
            response = self.upload(source)

        logged.assert_called_once()
        self.assertEqual(response.status_code, 400)
        self.assertIn("기준 엑셀", response.get_data(as_text=True))
        self.assertEqual(self.reference.read_bytes(), previous)
        self.assertIsNone(get_setting(self.conn, "reference_original_name"))

    def test_reference_upload_removes_new_file_when_none_existed(self):
        self.reference.unlink()
        source = self.make_changed_reference()

        with patch(
            "playauto.web.routes.set_settings",
            side_effect=sqlite3.OperationalError("boom"),
        ), patch.object(self.app.logger, "exception"):
            response = self.upload(source)

        self.assertEqual(response.status_code, 400)
        self.assertFalse(self.reference.exists())

    def test_export_failure_flashes_instead_of_500(self):
        with patch(
            "playauto.web.routes.export_to_dir",
            side_effect=RuntimeError("boom"),
        ), patch.object(self.app.logger, "exception") as logged:
            response = post_with_csrf(
                self.client,
                "/exports",
                {"option": "기타재화[35]"},
            )

        logged.assert_called_once()

        self.assertEqual(response.status_code, 302)
        body = self.client.get("/exports").get_data(as_text=True)
        self.assertIn("엑셀을 생성하지 못했습니다", body)

    def test_export_uses_commission_settings_for_blank_inputs(self):
        set_setting(self.conn, "commission_gmarket", "20")
        set_setting(self.conn, "commission_smartstore", "7")
        set_setting(self.conn, "commission_coupang", "11")

        with patch("playauto.web.routes.export_to_dir") as export:
            response = post_with_csrf(
                self.client,
                "/exports",
                {
                    "option": "기타재화[35]",
                    "commission_gmarket": "",
                    "commission_smartstore": "",
                    "commission_coupang": "",
                },
            )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(
            export.call_args.kwargs["commission"],
            {"지마켓": 20.0, "스마트스토어": 7.0, "쿠팡": 11.0},
        )

    def test_download_allow_list_and_delete(self):
        result = export_to_dir(
            self.db_path,
            self.reference,
            self.root / "exports",
            option="기타재화[35]",
        )

        for kind in ("setting", "single", "general"):
            response = self.client.get(
                f"/exports/{result.export_id}/{kind}"
            )
            self.assertEqual(response.status_code, 200)
            response.close()

        self.assertEqual(
            self.client.get(f"/exports/{result.export_id}/other").status_code,
            404,
        )
        self.assertEqual(
            self.client.get(
                f"/exports/{result.export_id}/%2e%2e%5coutside"
            ).status_code,
            404,
        )

        response = post_with_csrf(
            self.client,
            f"/exports/{result.export_id}/delete",
            {},
        )
        self.assertEqual(response.status_code, 302)
        self.assertFalse(result.directory.exists())
        self.assertEqual(
            post_with_csrf(
                self.client,
                f"/exports/{result.export_id}/delete",
                {},
            ).status_code,
            404,
        )
        self.assertEqual(
            post_with_csrf(self.client, "/exports/%2e%2e/delete", {}).status_code,
            404,
        )

    def test_download_race_returns_404(self):
        result = export_to_dir(
            self.db_path,
            self.reference,
            self.root / "exports",
            option="기타재화[35]",
        )

        with patch(
            "playauto.web.routes.export_file",
            side_effect=OSError("gone"),
        ):
            response = self.client.get(
                f"/exports/{result.export_id}/setting"
            )

        self.assertEqual(response.status_code, 404)
