import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "vendor"))
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook, load_workbook

from playauto.xlsx import (
    build_final_workbooks,
    load_phase1_input,
    output_paths,
    save_final_workbooks,
    save_setting_copy,
)


def make_input(path):
    wb = Workbook()
    setting = wb.active
    setting.title = "setting"
    setting.append(
        [
            "카테고리코드",
            "템플릿코드정보",
            "판매자관리코드",
            "온라인 상품명",
            "시장가격",
            "마진율",
            "기본이미지",
            "상세설명",
            "키워드",
            "제조사",
            "브랜드",
            "모델명",
            "단위 가격 표시 여부",
            "표시 용량",
            "표시 단위",
            "개당 용량",
            "구성 방식",
            None,
            "옵션조합",
            "옵션",
            "옵션 바코드",
            "옵션 모델번호",
        ]
    )
    setting.append(
        [
            "C1",
            "세트A",
            "MJ-0001",
            "상품1",
            10000,
            0.15,
            "https://img/1.jpg",
            None,
            "kw",
            "제조사",
            "브랜드",
            "모델",
            "N",
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
        ]
    )
    template = wb.create_sheet("template")
    template.append(
        [
            "쇼핑몰",
            "계정",
            "템플릿코드정보",
            "템플릿코드",
            "머리말/꼬리말 템플릿코드",
            "단일/일반",
        ]
    )
    template.append(["SmartStore", "A", "세트A", "111", "900", "일반"])
    template.append(["Gmarket", "B", "세트A", "222", "900", "일반"])
    template.append(["SmartStore", "A", "세트A", "333", "901", "단일"])
    wb.save(path)


class XlsxTests(unittest.TestCase):
    def test_load_phase1_input_requires_setting_and_template(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "input.xlsx"
            wb = Workbook()
            wb.active.title = "setting"
            wb.save(path)

            with self.assertRaisesRegex(ValueError, "template"):
                load_phase1_input(path)

    def test_output_paths_follow_original_names(self):
        single, general = output_paths("out", date(2026, 10, 27), False)

        self.assertEqual(single, Path("out/2026-10-27_플레이오토_단일.xlsx"))
        self.assertEqual(general, Path("out/2026-10-27_플레이오토_일반.xlsx"))
        self.assertEqual(
            output_paths("out", date(2026, 10, 27), True),
            (Path("out/2026-10-27_플레이오토_쿠팡.xlsx"), None),
        )

    def test_build_and_save_final_workbooks(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            path = tmp / "input.xlsx"
            make_input(path)
            wb, setting, template = load_phase1_input(path)

            built = build_final_workbooks(
                setting,
                template,
                selected_option="기타재화[35]",
                coupang_single=False,
                commission={"지마켓": 12.0, "스마트스토어": 6.0, "쿠팡": 10.0},
            )
            single_wb, single_ws = built.single_wb, built.single_ws
            general_wb, general_ws = built.general_wb, built.general_ws

            self.assertEqual(single_ws["A1"].value, "쇼핑몰(계정)")
            self.assertEqual(general_ws["A1"].value, "쇼핑몰(계정)")
            headers = [
                single_ws.cell(row=1, column=c).value
                for c in range(1, single_ws.max_column + 1)
            ]
            self.assertNotIn("옵션", headers)
            self.assertEqual(single_ws["A2"].value, "SmartStore=A")
            self.assertEqual(general_ws["A2"].value, "SmartStore=A\nGmarket=B")
            self.assertEqual(general_ws.max_row, 2)

            single_path, general_path = output_paths(
                tmp / "out", date(2026, 10, 27), False
            )
            save_final_workbooks(
                single_wb,
                general_wb,
                single_path,
                general_path,
            )
            self.assertTrue(single_path.exists())
            self.assertTrue(general_path.exists())

            saved = load_workbook(general_path, data_only=True)
            self.assertEqual(saved.active["A2"].value, "SmartStore=A\nGmarket=B")

    def test_save_setting_copy(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            path = tmp / "input.xlsx"
            make_input(path)

            copied = save_setting_copy(path, tmp / "out", date(2026, 10, 27))

            self.assertTrue(copied.exists())
            self.assertEqual(copied.name, "2026-10-27_플레이오토_설정.xlsx")
            self.assertEqual(
                load_workbook(copied, data_only=True).sheetnames,
                ["setting", "template"],
            )


if __name__ == "__main__":
    unittest.main()
