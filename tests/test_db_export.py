import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "vendor"))
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from playauto.constants import build_setting_headers
from playauto.db import ExportRow
from playauto.db_export import write_setting_from_rows


def make_reference(path):
    wb = Workbook()
    setting = wb.active
    setting.title = "setting"
    setting.append(["old"])
    setting.append(["old-row"])
    template = wb.create_sheet("template")
    template.append([
        "쇼핑몰", "계정", "템플릿코드정보", "템플릿코드",
        "머리말/꼬리말 템플릿코드", "단일/일반",
    ])
    template.append(["SmartStore", "A", "1박스", "111", "900", "일반"])
    template.append(["Gmarket", "B", "1박스", "222", "900", "일반"])
    template.append(["SmartStore", "A", "1박스", "333", "901", "단일"])
    template["A1"].font = Font(bold=True)
    sheet1 = wb.create_sheet("Sheet1")
    sheet1["A1"] = "보조"
    wb.save(path)


def make_export_row(**overrides):
    values = dict(
        listing_id=1, product_id=1, source_seq=1, seller_code="MJ-0001",
        sale_name="상품1", market_price=7080, margin_rate=0.15,
        image_url="https://img/1.jpg", order_url="https://order/1",
        category_code="C1", detail="상세", keywords="kw", manufacturer="제조",
        brand="브랜드", model="모델", unit_price_visible="Y",
        display_capacity="500", display_unit="ml", pack_config="묶음",
        capacity_per_unit="250", bundle_qty=1,
    )
    values.update(overrides)
    return ExportRow(**values)


class SettingWriterTests(unittest.TestCase):
    def test_write_setting_writes_headers_values_and_order_link(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            reference = tmp / "reference.xlsx"
            output = tmp / "out" / "setting.xlsx"
            make_reference(reference)

            result = write_setting_from_rows(reference, [make_export_row()], output)
            ws = load_workbook(result)["setting"]
            headers = [cell.value for cell in ws[1]]
            row = [cell.value for cell in ws[2]]

        self.assertEqual(len(headers), 18)
        self.assertEqual(headers, [*build_setting_headers()[:-1], "주문링크"])
        self.assertEqual(
            row,
            [
                "C1", "1박스", "MJ-0001", "상품1", 7080, 0.15,
                "https://img/1.jpg", "상세", "kw", "제조", "브랜드", "모델",
                "Y", "500", "ml", "250", "묶음", "https://order/1",
            ],
        )

    def test_write_setting_preserves_template_and_other_sheets(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            reference = tmp / "reference.xlsx"
            output = tmp / "setting.xlsx"
            make_reference(reference)
            before_wb = load_workbook(reference)
            before = [
                [cell.value for cell in row]
                for row in before_wb["template"].iter_rows()
            ]

            result = write_setting_from_rows(reference, [make_export_row()], output)
            after_wb = load_workbook(result)
            after = [
                [cell.value for cell in row]
                for row in after_wb["template"].iter_rows()
            ]

        self.assertEqual(after, before)
        self.assertTrue(after_wb["template"]["A1"].font.bold)
        self.assertEqual(after_wb["Sheet1"]["A1"].value, "보조")

    def test_write_setting_rejects_missing_sheets(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            reference = tmp / "reference.xlsx"
            wb = Workbook()
            wb.active.title = "setting"
            wb.save(reference)
            with self.assertRaisesRegex(ValueError, "template"):
                write_setting_from_rows(reference, [], tmp / "setting.xlsx")
