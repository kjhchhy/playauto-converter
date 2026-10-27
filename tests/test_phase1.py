import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "vendor"))
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook
from openpyxl.styles import Alignment

from playauto.phase1 import (
    apply_defaults,
    apply_template_codes,
    build_template_index,
    fill_sale_price_from_setting,
    get_header_map,
    populate_from_setting,
)


def make_sheet(rows):
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    return ws


def row_values(ws, row):
    return [ws.cell(row=row, column=c).value for c in range(1, ws.max_column + 1)]


class Phase1Tests(unittest.TestCase):
    def test_header_map_strips_text_and_skips_empty_cells(self):
        ws = make_sheet([["카테고리코드", None, " 판매가 ", ""], [1, 2, 3, 4]])

        self.assertEqual(
            get_header_map(ws), {"카테고리코드": 1, "판매가": 3, "": 4}
        )

    def test_template_index_groups_accounts_and_filters_mode(self):
        ws = make_sheet(
            [
                [
                    "쇼핑몰",
                    "계정",
                    "템플릿코드정보",
                    "템플릿코드",
                    "머리말/꼬리말 템플릿코드",
                    "단일/일반",
                ],
                ["SmartStore", "A", "세트A", "111", "900", "일반"],
                ["Gmarket", "B", "세트A", "222", "900", "일반"],
                ["SmartStore", "A", "세트A", "999", "999", "일반"],
                ["Auction", "C", "세트A", "333", "901", "단일"],
                ["11st", "D", "단일B", "444", "902", "단일"],
            ]
        )

        general = build_template_index(ws, "일반")
        single = build_template_index(ws, "단일")

        self.assertEqual(
            general["세트A"],
            {
                "mall_account": ["SmartStore=A", "Gmarket=B"],
                "tpl": ["111", "222"],
                "header_footer": ["900", "900"],
            },
        )
        self.assertEqual(set(single), {"세트A", "단일B"})
        self.assertEqual(single["세트A"]["mall_account"], ["Auction=C"])

    def test_template_index_rejects_missing_headers(self):
        ws = make_sheet([["쇼핑몰", "계정"], ["SmartStore", "A"]])

        with self.assertRaisesRegex(ValueError, "누락 헤더"):
            build_template_index(ws)

    def test_apply_template_codes_writes_multiline_and_blank_for_unknown(self):
        ws = make_sheet(
            [
                [
                    "쇼핑몰(계정)",
                    "템플릿코드",
                    "머리말/꼬리말 템플릿코드",
                    "템플릿코드정보",
                ],
                ["", "", "", "세트A"],
                ["old", "old", "old", "없음"],
            ]
        )
        index = {
            "세트A": {
                "mall_account": ["SmartStore=A", "Gmarket=B"],
                "tpl": ["111", "222"],
                "header_footer": ["900", "901"],
            }
        }

        apply_template_codes(ws, get_header_map(ws), index, "일반")

        self.assertEqual(row_values(ws, 2), ["SmartStore=A\nGmarket=B", "111\n222", "900\n901", "세트A"])
        self.assertEqual(row_values(ws, 3), ["", "", "", "없음"])
        self.assertEqual(
            ws.cell(row=2, column=1).alignment, Alignment(wrapText=True)
        )

    def test_apply_template_codes_requires_all_template_columns(self):
        ws = make_sheet([["템플릿코드정보", "템플릿코드"], ["세트A", "111"]])

        self.assertIsNone(apply_template_codes(ws, get_header_map(ws), {}))

    def test_populate_copies_fields_skips_blank_rows_and_unit_prices(self):
        src = make_sheet(
            [
                [
                    "카테고리코드",
                    "템플릿코드정보",
                    "판매자관리코드",
                    "온라인 상품명",
                    "시장가격",
                    "마진율",
                    "기본이미지",
                    "상세설명",
                    "키워드",
                    "제조사",
                    "브랜드",
                    "모델명",
                    "단위 가격 표시 여부",
                    "표시 용량",
                    "표시 단위",
                    "개당 용량",
                    "구성 방식",
                    "옵션조합",
                    "옵션",
                    "옵션 바코드",
                    "옵션 모델번호",
                ],
                ["C1", "T1", "MJ-0001", "상품1", 10000, 0.15, "img", "<p>x</p>", "kw", "제조", "브랜드", "모델", "Y", "10", "g", "1", "세트", "옵션없음", "", "", ""],
                [None, None, None, None],
                ["C2", "T2", "MJ-0002", "상품2", 20000, 0.2, "img2", None, "kw2", "", "", "", "N", "20", "ml", "2", "박스", "옵션없음", "", "", ""],
            ]
        )
        target = make_sheet(
            [
                [
                    "카테고리코드",
                    "온라인 상품명",
                    "상세설명",
                    "단위 가격 표시 여부",
                    "표시 용량",
                    "표시 단위",
                    "구성 방식",
                    "개당 용량",
                ]
            ]
        )
        general = make_sheet(
            [["카테고리코드", "단위 가격 표시 여부", "표시 용량"]]
        )

        populate_from_setting(
            src,
            [target, general],
            general_ws=general,
            coupang_single=False,
            detail_transform=lambda value: value.upper(),
        )

        self.assertEqual(
            row_values(target, 2),
            ["C1", "상품1", "<P>X</P>", "Y", "10", "g", "세트", "1"],
        )
        self.assertEqual(
            row_values(target, 3),
            ["C2", "상품2", None, "", "", "", "", ""],
        )
        self.assertEqual(row_values(general, 2), ["C1", "", ""])
        self.assertEqual(general.max_row, 3)

    def test_populate_requires_source_headers(self):
        src = make_sheet([["카테고리코드"]])
        target = make_sheet([["카테고리코드"]])

        with self.assertRaisesRegex(ValueError, "누락 헤더"):
            populate_from_setting(src, [target])

    def test_defaults_fill_fixed_notice_and_template_and_price(self):
        ws = make_sheet(
            [
                [
                    "템플릿코드정보",
                    "쇼핑몰(계정)",
                    "템플릿코드",
                    "머리말/꼬리말 템플릿코드",
                    "판매수량",
                    "원산지",
                    "배송방법",
                    "배송비",
                    "인증유형",
                    "미성년자 구매",
                    "상품분류코드",
                    "과세여부",
                    "옵션조합",
                    "판매가",
                    "상품정보제공고시 1",
                    "상품정보제공고시 3",
                    "상품정보제공고시 4",
                    "상품정보제공고시 8",
                ],
                ["세트A"],
            ]
        )
        settings = make_sheet(
            [
                ["시장가격", "마진율"],
                [10000, "15%"],
            ]
        )
        index = {
            "세트A": {
                "mall_account": ["SmartStore=A"],
                "tpl": ["111"],
                "header_footer": ["900"],
            }
        }

        apply_defaults(
            ws,
            product_code="35",
            tax_status="과세",
            selected_option="기타재화[35]",
            coupang_single=False,
            template_index=index,
            mode="일반",
            ws_setting=settings,
            commission={"지마켓": 12.0},
        )

        values = {h: ws.cell(row=2, column=c).value for h, c in get_header_map(ws).items()}
        self.assertEqual(values["판매수량"], 999)
        self.assertEqual(values["원산지"], "기타=상세설명참조")
        self.assertEqual(values["배송방법"], "무료")
        self.assertEqual(values["배송비"], 0)
        self.assertEqual(values["인증유형"], 2)
        self.assertEqual(values["미성년자 구매"], "Y")
        self.assertEqual(values["상품분류코드"], "35")
        self.assertEqual(values["과세여부"], "과세")
        self.assertEqual(values["옵션조합"], "옵션없음")
        self.assertEqual(values["쇼핑몰(계정)"], "SmartStore=A")
        self.assertEqual(values["상품정보제공고시 1"], "상품상세설명참조")
        self.assertEqual(values["상품정보제공고시 3"], "상품상세설명참조")
        self.assertIsNotNone(values["판매가"])

    def test_defaults_special_food_notices(self):
        ws = make_sheet(
            [
                [
                    "판매수량",
                    "상품정보제공고시 4",
                    "상품정보제공고시 8",
                    "상품정보제공고시 10",
                ],
                [None],
            ]
        )

        apply_defaults(
            ws,
            product_code="21",
            tax_status="과세",
            selected_option="가공식품[21]",
            coupang_single=False,
            template_index={},
            mode="일반",
            ws_setting=None,
            commission={},
        )

        self.assertEqual(
            ws["B2"].value,
            "수시입출고로 인하여 정확한 소비기한 기재가 어려우나 여유있는 상품으로 배송되고 있습니다.",
        )
        self.assertEqual(ws["C2"].value, "N")
        self.assertEqual(ws["D2"].value, "N")

    def test_coupang_defaults_leave_option_combination_and_insert_option(self):
        ws = make_sheet([["옵션조합", "옵션", "옵션 바코드", "옵션 모델번호"], [None, None, None, None]])

        apply_defaults(
            ws,
            product_code="35",
            tax_status="과세",
            selected_option="기타재화[35]",
            coupang_single=True,
            template_index={},
            mode="단일",
            ws_setting=None,
            commission={},
        )

        self.assertIsNone(ws["A2"].value)

    def test_fill_sale_price_uses_target_mode_commission(self):
        settings = make_sheet([["시장가격", "마진율"], [7080, 0.15]])
        general = make_sheet([["판매가"], [None]])
        single = make_sheet([["판매가"], [None]])

        fill_sale_price_from_setting(
            general,
            settings,
            general_ws=general,
            coupang_single=False,
            option_commission={"스마트스토어": 6.0, "지마켓": 12.0, "쿠팡": 10.0},
        )
        fill_sale_price_from_setting(
            single,
            settings,
            general_ws=general,
            coupang_single=False,
            option_commission={"스마트스토어": 6.0, "지마켓": 12.0, "쿠팡": 10.0},
        )

        self.assertEqual(general["A2"].value, 8900)
        self.assertEqual(single["A2"].value, 9500)


if __name__ == "__main__":
    unittest.main()
