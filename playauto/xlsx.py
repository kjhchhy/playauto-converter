import shutil
from dataclasses import dataclass
from datetime import date as date_type
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .constants import build_output_headers, meta_for_option
from .phase1 import (
    apply_defaults,
    build_template_index,
    populate_from_setting,
)


@dataclass(frozen=True)
class Phase1Workbooks:
    single_wb: Workbook
    single_ws: object
    general_wb: Workbook
    general_ws: object


def load_phase1_input(path):
    wb = load_workbook(path, data_only=True)
    missing = [name for name in ("setting", "template") if name not in wb.sheetnames]
    if missing:
        raise ValueError(f"필수 시트 누락: {missing}")
    return wb, wb["setting"], wb["template"]


def output_paths(output_dir, run_date, coupang_single):
    directory = Path(output_dir)
    prefix = f"{run_date.isoformat()}_플레이오토"
    if coupang_single:
        return directory / f"{prefix}_쿠팡.xlsx", None
    return (
        directory / f"{prefix}_단일.xlsx",
        directory / f"{prefix}_일반.xlsx",
    )


def setting_copy_path(output_dir, run_date):
    return Path(output_dir) / f"{run_date.isoformat()}_플레이오토_설정.xlsx"


def save_setting_copy(input_path, output_dir, run_date):
    destination = setting_copy_path(output_dir, run_date)
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(input_path, destination)
    return destination


def build_final_workbooks(
    setting_ws,
    template_ws,
    selected_option,
    coupang_single=False,
    commission=None,
):
    meta = meta_for_option(selected_option)
    notice_count = meta.notice_count
    tax_status = meta.tax_status
    product_code = meta.product_code
    commission = commission or {}

    single_wb = Workbook()
    single_ws = single_wb.active
    single_ws.title = "Sheet1"
    single_ws.append(
        build_output_headers("single", notice_count, coupang_single)
    )

    if coupang_single:
        populate_from_setting(
            setting_ws,
            [single_ws],
            general_ws=None,
            coupang_single=True,
        )
        apply_defaults(
            single_ws,
            product_code=product_code,
            tax_status=tax_status,
            selected_option=selected_option,
            coupang_single=True,
            template_index=build_template_index(template_ws, "단일"),
            mode="단일",
            ws_setting=setting_ws,
            general_ws=None,
            commission=commission,
        )
        return Phase1Workbooks(single_wb, single_ws, None, None)

    general_wb = Workbook()
    general_ws = general_wb.active
    general_ws.title = "Sheet1"
    general_ws.append(build_output_headers("general", notice_count, False))

    populate_from_setting(
        setting_ws,
        [single_ws, general_ws],
        general_ws=general_ws,
        coupang_single=False,
    )
    apply_defaults(
        single_ws,
        product_code=product_code,
        tax_status=tax_status,
        selected_option=selected_option,
        coupang_single=False,
        template_index=build_template_index(template_ws, "단일"),
        mode="단일",
        ws_setting=setting_ws,
        general_ws=general_ws,
        commission=commission,
    )
    apply_defaults(
        general_ws,
        product_code=product_code,
        tax_status=tax_status,
        selected_option=selected_option,
        coupang_single=False,
        template_index=build_template_index(template_ws, "일반"),
        mode="일반",
        ws_setting=setting_ws,
        general_ws=general_ws,
        commission=commission,
    )
    return Phase1Workbooks(single_wb, single_ws, general_wb, general_ws)


def save_final_workbooks(single_wb, general_wb, single_path, general_path):
    for workbook, destination in (
        (single_wb, single_path),
        (general_wb, general_path),
    ):
        if workbook is None or destination is None:
            continue
        destination = Path(destination)
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(destination)


def convert_phase1(
    input_path,
    output_dir,
    selected_option,
    coupang_single=False,
    commission=None,
    run_date=None,
    copy_setting=True,
):
    run_date = run_date or date_type.today()
    _, setting_ws, template_ws = load_phase1_input(input_path)
    workbooks = build_final_workbooks(
        setting_ws,
        template_ws,
        selected_option=selected_option,
        coupang_single=coupang_single,
        commission=commission,
    )
    single_path, general_path = output_paths(output_dir, run_date, coupang_single)
    save_final_workbooks(
        workbooks.single_wb,
        workbooks.general_wb,
        single_path,
        general_path,
    )
    setting_result = (
        save_setting_copy(input_path, output_dir, run_date)
        if copy_setting
        else None
    )
    return {
        "setting": setting_result,
        "single": single_path,
        "general": general_path,
    }
