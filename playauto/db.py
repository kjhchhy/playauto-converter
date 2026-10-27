import sqlite3
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from .pricing import parse_margin_ratio

DEFAULT_DB_PATH = Path("data/playauto.db")
SCHEMA_VERSION = 1

REQUIRED_HEADERS = ("순번", "대표이미지", "상품명", "판매가격", "URL")
OPTIONAL_HEADERS = (
    "판매자관리코드",
    "마진율",
    "상세설명",
    "키워드",
    "제조사",
    "브랜드",
    "모델명",
    "단위 가격 표시 여부",
    "표시 용량",
    "표시 단위",
    "개당 용량",
    "구성 방식",
)

SCHEMA_SQL = """
CREATE TABLE IF NOT EXISTS import_batches (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    imported_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    crawl_path TEXT,
    row_count INTEGER NOT NULL DEFAULT 0,
    inserted_count INTEGER NOT NULL DEFAULT 0,
    updated_count INTEGER NOT NULL DEFAULT 0,
    skipped_count INTEGER NOT NULL DEFAULT 0
);

CREATE TABLE IF NOT EXISTS products (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    source_seq INTEGER,
    name TEXT NOT NULL,
    image_url TEXT,
    order_url TEXT UNIQUE,
    category_code TEXT,
    detail TEXT,
    keywords TEXT,
    manufacturer TEXT,
    brand TEXT,
    model TEXT,
    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    last_seen_at TEXT,
    last_batch_id INTEGER REFERENCES import_batches(id)
);

CREATE TABLE IF NOT EXISTS listings (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    product_id INTEGER NOT NULL REFERENCES products(id),
    seller_code TEXT UNIQUE,
    bundle_qty INTEGER NOT NULL DEFAULT 1,
    sale_name TEXT,
    market_price INTEGER,
    margin_rate REAL DEFAULT 0.15
        CHECK (margin_rate >= 0 AND margin_rate < 1),
    unit_price_visible TEXT,
    display_capacity TEXT,
    display_unit TEXT,
    pack_config TEXT,
    capacity_per_unit TEXT,
    active INTEGER NOT NULL DEFAULT 1,
    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    UNIQUE(product_id, bundle_qty)
);

CREATE TABLE IF NOT EXISTS price_history (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    product_id INTEGER NOT NULL REFERENCES products(id),
    listing_id INTEGER REFERENCES listings(id),
    price INTEGER NOT NULL,
    observed_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    batch_id INTEGER REFERENCES import_batches(id),
    crawl_path TEXT
);

CREATE INDEX IF NOT EXISTS idx_listings_active ON listings(active);
CREATE INDEX IF NOT EXISTS idx_listings_product_id ON listings(product_id);
CREATE INDEX IF NOT EXISTS idx_products_source_seq ON products(source_seq);
"""


@dataclass(frozen=True)
class CrawlRow:
    source_seq: int
    name: str
    market_price: int
    image_url: str
    order_url: str
    seller_code: str | None = None
    margin_rate: float | None = None
    detail: str | None = None
    keywords: str | None = None
    manufacturer: str | None = None
    brand: str | None = None
    model: str | None = None
    unit_price_visible: str | None = None
    display_capacity: str | None = None
    display_unit: str | None = None
    pack_config: str | None = None
    capacity_per_unit: str | None = None


@dataclass(frozen=True)
class CrawlData:
    rows: tuple[CrawlRow, ...]
    skipped: int


@dataclass(frozen=True)
class ImportSummary:
    batch_id: int
    inserted: int
    updated: int
    skipped: int
    row_count: int


@dataclass(frozen=True)
class ExportRow:
    listing_id: int
    product_id: int
    source_seq: int
    seller_code: str
    sale_name: str
    market_price: int
    margin_rate: float
    image_url: str
    order_url: str
    category_code: str
    detail: str | None
    keywords: str | None
    manufacturer: str | None
    brand: str | None
    model: str | None
    unit_price_visible: str | None
    display_capacity: str | None
    display_unit: str | None
    pack_config: str | None
    capacity_per_unit: str | None
    bundle_qty: int


def connect(db_path: str | Path) -> sqlite3.Connection:
    path = Path(db_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(path))
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys=ON")
    return conn


def ensure_schema(conn: sqlite3.Connection) -> None:
    conn.executescript(SCHEMA_SQL)
    conn.execute(f"PRAGMA user_version={SCHEMA_VERSION}")
    conn.commit()


def parse_price(value) -> int:
    if value is None:
        raise ValueError("가격이 비어 있습니다.")
    text = str(value).strip().replace(",", "").replace(" ", "")
    if text.endswith("원"):
        text = text[:-1].strip()
    if not text:
        raise ValueError("가격이 비어 있습니다.")
    try:
        return int(float(text))
    except (TypeError, ValueError) as exc:
        raise ValueError(f"가격 형식이 올바르지 않습니다: {value!r}") from exc


def _is_blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _optional_text(value):
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        return text or None
    return str(value)


def _parse_margin(value):
    rate = parse_margin_ratio(value)
    if rate is None:
        return None
    if not 0 <= rate < 1:
        raise ValueError(f"마진율 범위가 올바르지 않습니다: {value!r}")
    return rate


def _wrap_parse(row_index, header, parser, value):
    try:
        return parser(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(
            f"[크롤] {row_index}행 '{header}' 값이 올바르지 않습니다."
        ) from exc


def _parse_sequence(value) -> int:
    text = str(value).strip().replace(",", "")
    if not text:
        raise ValueError("순번이 비어 있습니다.")
    number = float(text)
    if not number.is_integer():
        raise ValueError(f"순번 형식이 올바르지 않습니다: {value!r}")
    return int(number)


def load_crawl_rows(path: str | Path) -> CrawlData:
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.worksheets[0]
        row_iterator = ws.iter_rows(values_only=True)

        try:
            header_values = next(row_iterator)
        except StopIteration:
            header_values = ()

        header_map: dict[str, int] = {}
        for column, value in enumerate(header_values, start=1):
            if value is None:
                continue
            name = str(value).strip()
            if name and name not in header_map:
                header_map[name] = column

        missing = [
            header for header in REQUIRED_HEADERS if header not in header_map
        ]
        if missing:
            raise ValueError(f"누락 헤더: {', '.join(missing)}")

        rows: list[CrawlRow] = []
        skipped = 0
        for row_index, values in enumerate(row_iterator, start=2):
            if all(_is_blank(value) for value in values):
                skipped += 1
                continue

            def cell(header):
                column = header_map[header]
                if column > len(values):
                    return None
                return values[column - 1]

            for header in REQUIRED_HEADERS:
                if _is_blank(cell(header)):
                    raise ValueError(
                        f"[크롤] {row_index}행 '{header}' 값이 올바르지 않습니다."
                    )

            source_seq = _wrap_parse(
                row_index, "순번", _parse_sequence, cell("순번")
            )
            market_price = _wrap_parse(
                row_index, "판매가격", parse_price, cell("판매가격")
            )
            margin_rate = None
            if "마진율" in header_map:
                margin_rate = _wrap_parse(
                    row_index, "마진율", _parse_margin, cell("마진율")
                )

            def optional(header):
                if header not in header_map:
                    return None
                return _optional_text(cell(header))

            rows.append(
                CrawlRow(
                    source_seq=source_seq,
                    name=str(cell("상품명")).strip(),
                    market_price=market_price,
                    image_url=str(cell("대표이미지")).strip(),
                    order_url=str(cell("URL")).strip(),
                    seller_code=optional("판매자관리코드"),
                    margin_rate=margin_rate,
                    detail=optional("상세설명"),
                    keywords=optional("키워드"),
                    manufacturer=optional("제조사"),
                    brand=optional("브랜드"),
                    model=optional("모델명"),
                    unit_price_visible=optional("단위 가격 표시 여부"),
                    display_capacity=optional("표시 용량"),
                    display_unit=optional("표시 단위"),
                    pack_config=optional("구성 방식"),
                    capacity_per_unit=optional("개당 용량"),
                )
            )

        return CrawlData(rows=tuple(rows), skipped=skipped)
    finally:
        wb.close()


def _seller_code_taken(conn: sqlite3.Connection, code: str) -> bool:
    row = conn.execute(
        "SELECT 1 FROM listings WHERE seller_code = ?", (code,)
    ).fetchone()
    return row is not None


def _generate_seller_code(conn: sqlite3.Connection, source_seq: int) -> str:
    base = f"MJ-{source_seq:04d}"
    code = base
    suffix = 2
    while _seller_code_taken(conn, code):
        code = f"{base}-{suffix}"
        suffix += 1
    return code


def import_crawl(
    conn: sqlite3.Connection,
    crawl: CrawlData,
    category_code: str,
    margin: float,
    crawl_path: str,
) -> ImportSummary:
    if not category_code or not str(category_code).strip():
        raise ValueError("카테고리 코드가 비어 있습니다.")
    if margin is None or not 0 <= margin < 1:
        raise ValueError(f"마진율 범위가 올바르지 않습니다: {margin!r}")

    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    inserted = 0
    updated = 0

    with conn:
        cursor = conn.execute(
            "INSERT INTO import_batches (crawl_path) VALUES (?)", (crawl_path,)
        )
        batch_id = cursor.lastrowid

        for row in crawl.rows:
            existing_product = conn.execute(
                "SELECT id FROM products WHERE order_url = ?", (row.order_url,)
            ).fetchone()

            if existing_product is None:
                cursor = conn.execute(
                    "INSERT INTO products ("
                    "source_seq, name, image_url, order_url, category_code, "
                    "detail, keywords, manufacturer, brand, model, "
                    "created_at, updated_at, last_seen_at, last_batch_id"
                    ") VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                    (
                        row.source_seq,
                        row.name,
                        row.image_url,
                        row.order_url,
                        category_code,
                        row.detail,
                        row.keywords,
                        row.manufacturer,
                        row.brand,
                        row.model,
                        now,
                        now,
                        now,
                        batch_id,
                    ),
                )
                product_id = cursor.lastrowid
                inserted += 1
            else:
                product_id = existing_product["id"]
                conn.execute(
                    "UPDATE products SET "
                    "source_seq = ?, name = ?, image_url = ?, "
                    "last_seen_at = ?, last_batch_id = ?, updated_at = ? "
                    "WHERE id = ?",
                    (
                        row.source_seq,
                        row.name,
                        row.image_url,
                        now,
                        batch_id,
                        now,
                        product_id,
                    ),
                )
                updated += 1

            existing_listing = conn.execute(
                "SELECT id FROM listings WHERE product_id = ? AND bundle_qty = 1",
                (product_id,),
            ).fetchone()

            if existing_listing is None:
                if row.seller_code:
                    if _seller_code_taken(conn, row.seller_code):
                        raise ValueError(
                            f"[크롤] {row.source_seq}행 판매자관리코드 중복: "
                            f"{row.seller_code}"
                        )
                    seller_code = row.seller_code
                else:
                    seller_code = _generate_seller_code(conn, row.source_seq)

                margin_rate = (
                    row.margin_rate if row.margin_rate is not None else margin
                )
                cursor = conn.execute(
                    "INSERT INTO listings ("
                    "product_id, seller_code, bundle_qty, sale_name, market_price, "
                    "margin_rate, unit_price_visible, display_capacity, "
                    "display_unit, pack_config, capacity_per_unit, active, "
                    "created_at, updated_at"
                    ") VALUES (?, ?, 1, ?, ?, ?, ?, ?, ?, ?, ?, 1, ?, ?)",
                    (
                        product_id,
                        seller_code,
                        row.name,
                        row.market_price,
                        margin_rate,
                        row.unit_price_visible,
                        row.display_capacity,
                        row.display_unit,
                        row.pack_config,
                        row.capacity_per_unit,
                        now,
                        now,
                    ),
                )
                listing_id = cursor.lastrowid
            else:
                listing_id = existing_listing["id"]
                conn.execute(
                    "UPDATE listings SET market_price = ?, updated_at = ? "
                    "WHERE id = ?",
                    (row.market_price, now, listing_id),
                )

            conn.execute(
                "INSERT INTO price_history ("
                "product_id, listing_id, price, observed_at, batch_id, crawl_path"
                ") VALUES (?, ?, ?, ?, ?, ?)",
                (
                    product_id,
                    listing_id,
                    row.market_price,
                    now,
                    batch_id,
                    crawl_path,
                ),
            )

        row_count = len(crawl.rows)
        conn.execute(
            "UPDATE import_batches SET "
            "row_count = ?, inserted_count = ?, updated_count = ?, skipped_count = ? "
            "WHERE id = ?",
            (row_count, inserted, updated, crawl.skipped, batch_id),
        )

    return ImportSummary(
        batch_id=batch_id,
        inserted=inserted,
        updated=updated,
        skipped=crawl.skipped,
        row_count=row_count,
    )
