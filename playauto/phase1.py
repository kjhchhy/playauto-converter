from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from .constants import (
    COUPANG_OPTION_HEADERS,
    UNIT_PRICE_FLAG_OUTPUT,
    UNIT_PRICE_FLAG_SOURCE,
    UNIT_PRICE_SOURCE_HEADERS,
)
from .pricing import (
    calculate_base_payout,
    ceil_to_100,
    parse_margin_ratio,
)

TEMPLATE_REQUIRED_HEADERS = [
    "쇼핑몰",
    "계정",
    "템플릿코드정보",
    "템플릿코드",
    "머리말/꼬리말 템플릿코드",
    "단일/일반",
]

TEMPLATE_OUTPUT_HEADERS = [
    "템플릿코드정보",
    "쇼핑몰(계정)",
    "템플릿코드",
    "머리말/꼬리말 템플릿코드",
]

SOURCE_HEADERS = [
    "카테고리코드",
    "템플릿코드정보",
    "판매자관리코드",
    "온라인 상품명",
    "시장가격",
    "마진율",
    "기본이미지",
    "상세설명",
    "키워드",
    "제조사",
    "브랜드",
    "모델명",
]

FIXED_DEFAULTS = {
    "판매수량": 999,
    "원산지": "기타=상세설명참조",
    "배송방법": "무료",
    "배송비": 0,
    "인증유형": 2,
    "미성년자 구매": "Y",
}

SPECIAL_FOOD_OPTIONS = {"가공식품[21]", "가공식품_유아동[21]"}
HEALTH_FOOD_OPTION = "건강기능식품[22]"

FOOD_NOTICE_4 = (
    "수시입출고로 인하여 정확한 소비기한 기재가 어려우나 "
    "여유있는 상품으로 배송되고 있습니다."
)
HEALTH_NOTICE_3 = (
    FOOD_NOTICE_4
    + " (자세한 문의는 문의글 혹은 판매자 고객센터로 문의해주세요.)"
)


def get_header_map(ws):
    if ws is None or ws.max_row < 1:
        return {}
    return {
        str(cell.value).strip(): cell.column
        for cell in ws[1]
        if cell.value is not None
    }


def build_template_index(ws_template, mode_filter=None):
    if ws_template is None or ws_template.max_row < 2:
        return {}
    header_map = get_header_map(ws_template)
    missing = [header for header in TEMPLATE_REQUIRED_HEADERS if header not in header_map]
    if missing:
        raise ValueError(f"[template] 누락 헤더: {missing}")

    index = {}
    seen = {}
    for row in range(2, ws_template.max_row + 1):
        def text(header):
            value = ws_template.cell(row=row, column=header_map[header]).value
            return str(value).strip() if value is not None else ""

        mode_value = text("단일/일반")
        if mode_filter and mode_value != mode_filter:
            continue
        key = text("템플릿코드정보")
        if not key:
            continue

        if key not in index:
            index[key] = {"mall_account": [], "tpl": [], "header_footer": []}
            seen[key] = set()

        pair = (text("쇼핑몰"), text("계정"))
        if pair in seen[key]:
            continue
        seen[key].add(pair)
        index[key]["mall_account"].append(f"{pair[0]}={pair[1]}")
        index[key]["tpl"].append(text("템플릿코드"))
        index[key]["header_footer"].append(text("머리말/꼬리말 템플릿코드"))
    return index


def apply_template_codes(ws, header_map, template_index, mode=None):
    if any(header not in header_map for header in TEMPLATE_OUTPUT_HEADERS):
        return None

    col_info = header_map["템플릿코드정보"]
    col_mall_account = header_map["쇼핑몰(계정)"]
    col_tpl = header_map["템플릿코드"]
    col_header_footer = header_map["머리말/꼬리말 템플릿코드"]

    for row in range(2, ws.max_row + 1):
        value = ws.cell(row=row, column=col_info).value
        key = str(value).strip() if value is not None else ""
        if not key:
            continue

        data = template_index.get(key)
        if not data:
            for column in (col_mall_account, col_tpl, col_header_footer):
                ws.cell(row=row, column=column).value = ""
            continue

        ws.cell(row=row, column=col_mall_account).value = "\n".join(
            data["mall_account"]
        )
        ws.cell(row=row, column=col_tpl).value = "\n".join(data["tpl"])
        ws.cell(row=row, column=col_header_footer).value = "\n".join(
            data["header_footer"]
        )
        for column in (col_mall_account, col_tpl, col_header_footer):
            ws.cell(row=row, column=column).alignment = Alignment(wrapText=True)

    letter = get_column_letter(col_mall_account)
    dimension = ws.column_dimensions[letter]
    dimension.width = (dimension.width or 8.43) * 2
    return None


def populate_from_setting(
    ws_setting,
    targets,
    general_ws=None,
    coupang_single=False,
    detail_transform=None,
):
    src_headers_all = [*SOURCE_HEADERS, *UNIT_PRICE_SOURCE_HEADERS]
    if coupang_single:
        src_headers_all += ["옵션조합", "옵션", *COUPANG_OPTION_HEADERS]

    excluded = {"시장가격", "마진율", *UNIT_PRICE_SOURCE_HEADERS}
    required_src = [header for header in src_headers_all if header not in excluded]

    src_map = get_header_map(ws_setting)
    missing_src = [header for header in required_src if header not in src_map]
    if missing_src:
        raise ValueError(f"[원본(setting)] 누락 헤더: {missing_src}")

    dst_maps = {ws: get_header_map(ws) for ws in targets}

    for row in range(2, ws_setting.max_row + 1):
        if not any(
            ws_setting.cell(row=row, column=column).value
            for column in range(1, ws_setting.max_column + 1)
        ):
            continue

        row_buffers = {
            ws: ["" for _ in range(ws.max_column)] for ws in targets
        }
        unit_price_flag = ""
        if UNIT_PRICE_FLAG_SOURCE in src_map:
            value = ws_setting.cell(
                row=row, column=src_map[UNIT_PRICE_FLAG_SOURCE]
            ).value
            unit_price_flag = str(value or "").strip().upper()

        for header in src_headers_all:
            if header not in src_map:
                continue
            if header in UNIT_PRICE_SOURCE_HEADERS and unit_price_flag != "Y":
                continue

            value = ws_setting.cell(row=row, column=src_map[header]).value
            if header == "상세설명" and isinstance(value, str) and detail_transform:
                value = detail_transform(value)
            target_header = (
                UNIT_PRICE_FLAG_OUTPUT
                if header == UNIT_PRICE_FLAG_SOURCE
                else header
            )

            for ws in targets:
                if header in UNIT_PRICE_SOURCE_HEADERS and ws is general_ws:
                    continue
                dst_map = dst_maps[ws]
                if target_header not in dst_map:
                    continue
                row_buffers[ws][dst_map[target_header] - 1] = value

        for ws, buffer in row_buffers.items():
            ws.append(buffer)
    return None


def fill_sale_price_from_setting(
    ws_target,
    ws_setting,
    general_ws=None,
    coupang_single=False,
    option_commission=None,
):
    if ws_target is None or ws_target.max_row < 2:
        return None

    dst_map = get_header_map(ws_target)
    col_sale = dst_map.get("판매가")
    src_map = {}
    if ws_setting is not None and ws_setting.max_row >= 2:
        src_map = get_header_map(ws_setting)

    use_setting = "시장가격" in src_map and "마진율" in src_map
    option_commission = option_commission or {}
    if coupang_single:
        pct_key = "쿠팡"
    elif ws_target is general_ws:
        pct_key = "스마트스토어"
    else:
        pct_key = "지마켓"
    pct = option_commission.get(pct_key) or 0.0

    for row in range(2, ws_target.max_row + 1):
        if use_setting and row <= ws_setting.max_row:
            market_price = ws_setting.cell(
                row=row, column=src_map["시장가격"]
            ).value
            margin_rate = ws_setting.cell(row=row, column=src_map["마진율"]).value
        else:
            market_price = (
                ws_target.cell(row=row, column=dst_map["시장가격"]).value
                if "시장가격" in dst_map
                else None
            )
            margin_rate = (
                ws_target.cell(row=row, column=dst_map["마진율"]).value
                if "마진율" in dst_map
                else None
            )

        base_payout = calculate_base_payout(market_price, parse_margin_ratio(margin_rate))
        if base_payout is None or not col_sale:
            continue

        final_price = ceil_to_100(base_payout / (1.0 - float(pct) / 100.0))
        ws_target.cell(row=row, column=col_sale).value = final_price
    return None


def apply_defaults(
    ws,
    product_code,
    tax_status,
    selected_option,
    coupang_single=False,
    template_index=None,
    mode=None,
    ws_setting=None,
    general_ws=None,
    commission=None,
):
    if ws is None or ws.max_row < 2:
        return None

    header_map = get_header_map(ws)
    max_row = ws.max_row

    fixed_values = {
        **FIXED_DEFAULTS,
        "상품분류코드": product_code,
        "과세여부": tax_status,
    }
    if not coupang_single:
        fixed_values["옵션조합"] = "옵션없음"

    for header, fixed in fixed_values.items():
        column = header_map.get(header)
        if not column:
            continue
        for row in range(2, max_row + 1):
            ws.cell(row=row, column=column).value = fixed

    option = str(selected_option).strip()
    is_special_food = option in SPECIAL_FOOD_OPTIONS
    is_health_food = option == HEALTH_FOOD_OPTION

    for header, column in header_map.items():
        if "상품정보제공고시" not in header:
            continue
        default_value = "상품상세설명참조"
        clean_header = header.replace("  ", " ").strip()
        if is_special_food:
            if clean_header == "상품정보제공고시 4":
                default_value = FOOD_NOTICE_4
            elif clean_header in ("상품정보제공고시 8", "상품정보제공고시 10"):
                default_value = "N"
        elif is_health_food and clean_header == "상품정보제공고시 3":
            default_value = HEALTH_NOTICE_3

        for row in range(2, max_row + 1):
            ws.cell(row=row, column=column).value = default_value

    apply_template_codes(ws, header_map, template_index or {}, mode)
    fill_sale_price_from_setting(
        ws,
        ws_setting,
        general_ws=general_ws,
        coupang_single=coupang_single,
        option_commission=commission,
    )
    return None
