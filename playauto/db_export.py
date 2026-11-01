from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from .constants import COUPANG_OPTION_HEADERS, build_setting_headers
from .db import (
    ExportRow, connect, ensure_schema, fetch_export_rows, validate_export_rows,
)
from .xlsx import convert_phase1, setting_copy_path

REQUIRED_SHEETS = ("setting", "template")
TEMPLATE_INFO = "1박스"
ORDER_URL_HEADER = "주문링크"


OPTION_HEADERS = ["옵션조합", "옵션", *COUPANG_OPTION_HEADERS]


def _setting_headers(coupang_single=False):
    headers = [*build_setting_headers()[:-1], ORDER_URL_HEADER]
    if coupang_single:
        headers[-1:-1] = OPTION_HEADERS
    return headers


def _row_values(row: ExportRow, coupang_single=False):
    values = [
        row.category_code,
        TEMPLATE_INFO,
        row.seller_code,
        row.sale_name,
        row.market_price,
        row.margin_rate,
        row.image_url,
        row.detail_page_url or row.detail,
        row.keywords,
        row.manufacturer,
        row.brand,
        row.model,
        row.unit_price_visible,
        row.display_capacity,
        row.display_unit,
        row.capacity_per_unit,
        row.pack_config,
        row.order_url,
    ]
    if coupang_single:
        values[-1:-1] = [None] * len(OPTION_HEADERS)
    return values


def write_setting_from_rows(
    reference_path: str | Path,
    rows: list[ExportRow],
    output_path: str | Path,
    coupang_single=False,
) -> Path:
    wb = load_workbook(reference_path)
    try:
        missing = [name for name in REQUIRED_SHEETS if name not in wb.sheetnames]
        if missing:
            raise ValueError(f"필수 시트 누락: {missing}")

        ws = wb["setting"]
        ws.delete_rows(1, ws.max_row)
        ws.append(_setting_headers(coupang_single))
        for row in rows:
            ws.append(_row_values(row, coupang_single))

        destination = Path(output_path)
        destination.parent.mkdir(parents=True, exist_ok=True)
        wb.save(destination)
    finally:
        wb.close()
    return destination


def export_from_db(
    db_path,
    reference_path,
    output_dir,
    selected_option,
    coupang_single=False,
    commission=None,
    run_date=None,
) -> dict:
    run_date = run_date or date.today()

    conn = connect(db_path)
    try:
        ensure_schema(conn)
        rows = fetch_export_rows(conn)
        validate_export_rows(rows)
    finally:
        conn.close()

    setting_path = write_setting_from_rows(
        reference_path,
        rows,
        setting_copy_path(output_dir, run_date),
        coupang_single=coupang_single,
    )
    result = convert_phase1(
        setting_path,
        output_dir,
        selected_option,
        coupang_single,
        commission,
        run_date,
        copy_setting=False,
    )
    return {
        "setting": setting_path,
        "single": result["single"],
        "general": result["general"],
    }
