import hashlib
import os
import shutil
import uuid
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from playauto.db_export import export_from_db

REQUIRED_SHEETS = ("setting", "template")
EXPORT_KINDS = ("setting", "single", "general")
SETTING_SUFFIX = "_설정.xlsx"
SINGLE_SUFFIXES = ("_단일.xlsx", "_쿠팡.xlsx")
GENERAL_SUFFIX = "_일반.xlsx"


@dataclass(frozen=True)
class ExportResult:
    export_id: str
    directory: Path
    setting: Path
    single: Path
    general: Path | None


def validate_reference(path: str | Path) -> None:
    source = Path(path)
    if not source.is_file():
        raise ValueError("기준 엑셀 파일이 없습니다.")
    try:
        workbook = load_workbook(source, read_only=True)
    except Exception as exc:  # noqa: BLE001 - surface any decode failure
        raise ValueError("기준 엑셀 파일을 읽을 수 없습니다.") from exc
    try:
        missing = [
            name for name in REQUIRED_SHEETS if name not in workbook.sheetnames
        ]
    finally:
        workbook.close()
    if missing:
        raise ValueError(f"필수 시트 누락: {missing}")


def save_reference(src: str | Path, dest: str | Path) -> dict:
    source = Path(src)
    destination = Path(dest)
    validate_reference(source)
    destination.parent.mkdir(parents=True, exist_ok=True)

    temporary = destination.with_name(
        f"{destination.name}.tmp-{uuid.uuid4().hex}"
    )
    try:
        shutil.copyfile(source, temporary)
        os.replace(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)

    digest = hashlib.sha256(destination.read_bytes()).hexdigest()
    return {
        "original_name": source.name,
        "size": destination.stat().st_size,
        "sha256": digest,
    }


def _as_run_date(run_date):
    if isinstance(run_date, str):
        return date.fromisoformat(run_date)
    return run_date


def export_to_dir(
    db_path,
    reference_path,
    export_root,
    *,
    option,
    run_date=None,
    coupang_single=False,
    commission=None,
) -> ExportResult:
    reference = Path(reference_path)
    if not reference.is_file():
        raise ValueError("기준 엑셀 파일이 없습니다.")
    validate_reference(reference)

    export_id = (
        datetime.now().strftime("%Y%m%d-%H%M%S")
        + "-"
        + uuid.uuid4().hex[:8]
    )
    directory = Path(export_root) / export_id
    directory.mkdir(parents=True, exist_ok=True)

    try:
        result = export_from_db(
            db_path,
            reference,
            directory,
            option,
            coupang_single=coupang_single,
            commission=commission,
            run_date=_as_run_date(run_date),
        )
    except Exception:
        shutil.rmtree(directory, ignore_errors=True)
        raise

    general = result.get("general")
    return ExportResult(
        export_id=export_id,
        directory=directory,
        setting=Path(result["setting"]),
        single=Path(result["single"]),
        general=Path(general) if general else None,
    )


def export_kind(filename: str) -> str | None:
    if filename.endswith(SETTING_SUFFIX):
        return "setting"
    if filename.endswith(SINGLE_SUFFIXES):
        return "single"
    if filename.endswith(GENERAL_SUFFIX):
        return "general"
    return None


def export_file(directory: str | Path, kind: str) -> Path | None:
    if kind not in EXPORT_KINDS:
        return None
    folder = Path(directory)
    for entry in sorted(folder.iterdir()):
        if entry.is_file() and export_kind(entry.name) == kind:
            return entry
    return None


def list_exports(export_root) -> list[dict]:
    root = Path(export_root)
    if not root.is_dir():
        return []

    exports = []
    for directory in root.iterdir():
        if not directory.is_dir():
            continue
        files = []
        for entry in sorted(directory.iterdir()):
            if not entry.is_file():
                continue
            kind = export_kind(entry.name)
            if kind is None:
                continue
            stat = entry.stat()
            files.append(
                {
                    "name": entry.name,
                    "kind": kind,
                    "size": stat.st_size,
                    "modified_at": datetime.fromtimestamp(
                        stat.st_mtime
                    ).isoformat(timespec="seconds"),
                }
            )
        exports.append(
            {
                "export_id": directory.name,
                "created_at": datetime.fromtimestamp(
                    directory.stat().st_mtime
                ).isoformat(timespec="seconds"),
                "files": files,
            }
        )

    exports.sort(
        key=lambda item: (item["created_at"], item["export_id"]),
        reverse=True,
    )
    return exports


def _safe_export_id(export_id: str) -> str | None:
    if not export_id or export_id in (".", ".."):
        return None
    if "/" in export_id or "\\" in export_id or ".." in export_id:
        return None
    if Path(export_id).is_absolute():
        return None
    return export_id


def delete_export(export_root, export_id: str) -> bool:
    safe_id = _safe_export_id(export_id)
    if safe_id is None:
        return False

    root = Path(export_root)
    target = root / safe_id
    if not target.is_dir():
        return False
    try:
        if target.resolve().parent != root.resolve():
            return False
    except OSError:
        return False

    shutil.rmtree(target)
    return True
